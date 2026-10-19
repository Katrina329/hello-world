# -*- coding: utf-8 -*-
import shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn

SRC = r'D:\桌面\数据结构-实验二-指导书.docx'
DST = r'D:\桌面\ai\数据结构-实验二-指导书-已完成.docx'
shutil.copyfile(SRC, DST)
doc = Document(DST)

# ---------- helpers ----------
def clear_para(p):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)

def set_font(run, name=None, size=None, bold=None):
    if name:
        run.font.name = name
        rpr = run._element.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = rpr.makeelement(qn('w:rFonts'), {})
            rpr.append(rf)
        rf.set(qn('w:ascii'), name)
        rf.set(qn('w:hAnsi'), name)
        rf.set(qn('w:eastAsia'), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold

def set_para_text(p, text, font_name=None, size=None, bold=None):
    clear_para(p)
    lines = text.split('\n')
    run = p.add_run()
    for i, line in enumerate(lines):
        if i:
            run.add_break(WD_BREAK.LINE)
        run.add_text(line)
    set_font(run, font_name, size, bold)
    return run

def set_cell_text(cell, text, font_name=None, size=None, bold=None):
    p = cell.paragraphs[0]
    # remove any extra paragraphs beyond first
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    set_para_text(p, text, font_name, size, bold)

def insert_image(cell, path, width_in=6.0):
    p = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    clear_para(p)
    run = p.add_run()
    run.add_picture(path, width=Inches(width_in))

# ---------- content ----------
P0010_CODE = '''import sys


def main():
    data = sys.stdin.read().split()
    if not data:
        return
    n = int(data[0])
    a = list(map(int, data[1:1 + n]))
    b = list(map(int, data[1 + n:1 + 2 * n]))
    i = 0
    j = 0
    result = []
    while i < n and j < n:
        if a[i] <= b[j]:
            result.append(a[i])
            i += 1
        else:
            result.append(b[j])
            j += 1
    if i < n:
        result.extend(a[i:])
    if j < n:
        result.extend(b[j:])
    print(" ".join(map(str, result)))


if __name__ == "__main__":
    main()'''

P0030_CODE = '''def josephus(n, m):
    # 返回 n 只猴子、每次数到 m 出圈时最后剩下的猴子编号
    s = 0
    for i in range(2, n + 1):
        s = (s + m) % i
    return s + 1


def main():
    results = []
    while True:
        line = input()
        if not line.strip():
            continue
        n, m = map(int, line.split())
        if n == 0 and m == 0:
            break
        results.append(str(josephus(n, m)))
    print("\\n".join(results))


if __name__ == "__main__":
    main()'''

P0050_CODE = '''class Node:
    def __init__(self, data, next=None):
        self.data, self.next = data, next


class LinkList:
    def __init__(self, lst):
        self.head = Node(lst[0])
        p = self.head
        for i in lst[1:]:
            node = Node(i)
            p.next = node
            p = p.next


def genLinkList():
    a = list(map(int, input().split()))
    b = list(map(int, input().split()))
    n = int(input())
    la = LinkList(a)
    lb = LinkList(b)
    p = la.head
    last = None
    while p is not None:
        last = p
        p = p.next
    p = lb.head
    for i in range(n - 1):
        p = p.next
    last.next = p
    return la, lb


a, b = genLinkList()


def findJointPoint(a, b):
    len_a = 0
    p = a.head
    while p is not None:
        len_a += 1
        p = p.next
    len_b = 0
    p = b.head
    while p is not None:
        len_b += 1
        p = p.next
    pa, pb = a.head, b.head
    if len_a > len_b:
        for _ in range(len_a - len_b):
            pa = pa.next
    else:
        for _ in range(len_b - len_a):
            pb = pb.next
    while pa is not pb:
        pa = pa.next
        pb = pb.next
    return pa.data


print(findJointPoint(a, b))'''

def numbered(code):
    lines = code.split('\n')
    w = len(str(len(lines)))
    return '\n'.join((str(i+1).rjust(w) + '  ' + line) for i, line in enumerate(lines))

# paragraph fills by original index
par_fills = {
    98: '输入一个整数 n，再输入两个各含 n 个整数的升序序列，要求用 O(n) 时间把两个序列合并成一个新的升序序列并输出。',
    102: '题目本质是二路归并。输入两个长度均为 n 的非降序序列，需要输出长度为 2n 的合并结果。由于两个序列本身已经有序，每次只需比较两个序列当前剩余部分的首元素，把较小者放入结果，即可保证结果整体有序。每行整数不超过 100000 个，n 最大为 100000，因此必须在线性时间内完成，不能先把两个序列拼接起来再整体排序。',
    104: '选用双指针二路归并。指针 i 指向序列 a 的当前元素，指针 j 指向序列 b 的当前元素，result 保存结果。每轮比较 a[i] 与 b[j]，把较小者加入 result 并后移对应指针；当某一序列被取完后，把另一序列剩余部分一次性追加。选择该算法是因为两个序列本身有序，双指针每个元素只访问一次，时间严格为 O(n)，实现简单且稳定。',
    106: '初始时 i=j=0，result 为空。两个子序列分别有序，所以每轮取两个剩余首元素中较小者加入 result，既能保证 result 末尾不大于新加入元素，也能保证不遗漏元素；被移动的指针不会回头，因此每个元素恰好被访问一次。当一方被取完后，另一方剩余部分仍保持升序且都大于 result 当前末尾，直接 extend 即可。最终 result 长度为 2n，包含全部输入元素。',
    108: '时间复杂度 O(n)：每轮比较都会使 i 或 j 至少增加 1，两个序列共 2n 个元素，循环总次数不超过 2n。空间复杂度 O(n)：主要是保存结果的 result 列表，以及两个输入列表 a、b，均与 n 同阶。n 最大 100000，200000 次线性操作远小于 1 秒限制，不会超时。',
    110: '主要考虑了以下边界：两个序列中出现相等元素时，用 a[i] <= b[j] 优先取 a，保证相等元素也能稳定进入结果；某一序列先被取完后，剩余部分用切片一次性 extend，并用 i < n and j < n 防止下标越界；读入为空时直接 return，避免访问 data[0] 出错。',
    120: '一次通过。提交前先在本地运行样例，确认输出与样例一致；再用随机生成的多组升序序列（含重复元素和一方先取完的情况）与 Python 内置 sorted 合并排序结果逐一比对，全部一致后提交，平台评测结果为 Accepted。',
    125: '输入若干行，每行给出 n 和 m，表示 n 只猴子从 1 号起顺时针报数，数到 m 的出圈，求最后剩下的猴王编号；以一行 0 0 表示结束，每组输出一行。',
    129: '这是约瑟夫环问题。给定 n 和 m，每轮删除当前环中的第 m 个元素，删除 n-1 次后只剩一个元素，要求输出它的原始编号。输入含多组数据，以 n=m=0 结束，且 0<n,m<=300。虽然 n 不大，但题目要求思路清晰，采用递推公式可以在 O(n) 时间内直接计算每一组答案。',
    131: '采用约瑟夫问题的逆推公式。设 f(i) 表示 i 个人报数到 m 出列时，最后幸存者在从 0 开始编号的环中的位置，则 f(1)=0，f(i)=(f(i-1)+m) % i。循环从 i=2 递推到 i=n，得到 f(n) 后加 1 还原为从 1 开始的编号。选择递推而不是模拟删除，是因为递推每组只需 O(n) 时间和 O(1) 额外空间，代码简洁，也不易在报数边界处出错。',
    133: '递推的含义是：f(i-1) 是 i-1 个人时的幸存者位置；当人数从 i-1 增加到 i 时，可以看作上一次删除后从被删位置的下一位重新编号为 0，幸存者在新的 0 号坐标系中位置为 f(i-1)，把它映射回 i 个人的 0 号坐标系需要加上 m 并取模 i，即 (f(i-1)+m) % i。因此由 f(1)=0 递推到 n 即得 0 号坐标系下的幸存者编号，再加 1 得到题目要求的 1 号起点编号。',
    135: '时间复杂度 O(n)：每组数据循环 2 到 n 共 n-1 次常数运算；若有 k 组数据，总时间为 O(k·n)。空间复杂度 O(1)：递推只使用 s、i 两个变量；results 列表用于保存各组输出，是必要的输出存储。n,m 最大 300，任何一组都能在极短时间内完成，不会超时。',
    137: '主要考虑了：m 大于当前人数时取模运算使报数自然在环内循环，不会越界；n=1 时 for 循环不执行，直接返回 s+1=1；输入中出现空行时用 if not line.strip(): continue 跳过；以 0 0 作为结束标记，先判断结束再处理，避免把 0 0 当作一组数据。',
    147: '一次通过。先用样例 6 2、12 4、8 3 验证输出为 5、1、7，再手工计算 n=1 和 m>n 的几组小数据，确认递推结果与逐次删除模拟结果一致，最后在 OpenJudge 提交并得到 Accepted。',
    152: '给定两个单链表 a 和 b，程序先把 a 的尾结点连接到 b 的第 n 个结点，使两条链表形成 Y 形；要求补全 findJointPoint 函数，在线性时间内返回二者共享的第一个结点的数据。',
    156: '题目已经给出 Node、LinkList 与 genLinkList 框架。genLinkList 读入两个序列并令 a 的尾结点 next 指向 b 的第 n 个结点，使 a 和 b 共享从 b 的第 n 个结点开始的后缀。findJointPoint 需要返回分叉点。由于分叉点之后两条链表完全相同，从两个头结点到分叉点的路径长度不同，但共享后缀长度相同。',
    158: '采用长度对齐的双指针法。先分别遍历 a、b 求出长度 len_a、len_b；然后让较长链表的指针先走长度差 |len_a-len_b| 步，使两个指针到分叉点的剩余距离相等；之后两个指针同步前进，第一次指向同一结点即为共享首结点。选择该算法是因为它只用常数空间，总遍历次数为 O(len_a+len_b)，满足题目 O(n) 要求。',
    160: '设 a 的独占前缀长为 x，b 的独占前缀长为 y，共享后缀长为 z，则 len_a=x+z，len_b=y+z。较长的链表先走 |x-y| 步后，两个指针分别位于从各自头结点出发、到分叉点距离相同的位置；同步移动时，若尚未相遇则两指针不可能指向同一结点，因此可以安全地同时后移，不会跳过交点；到达分叉点时两者必然同时指向同一结点，循环结束，返回 pa.data。',
    162: '时间复杂度 O(n)：先遍历求两个链表长度共 O(len_a+len_b)，再移动指针至相遇最多 O(len_a+len_b)，整体为线性。空间复杂度 O(1)：只使用 pa、pb 和两个长度变量等常数空间。两个链表总长度最多 200000，线性操作可在 1 秒内完成，不会超时。',
    164: '考虑了以下边界：len_a 与 len_b 相等时不跳过、直接同步扫描；分叉点为 b 的第一个结点（n=1）时，a 先走完独占前缀后与 pb 在 b.head 相遇；分叉点位于链表末尾时共享后缀只有一个结点，while pa is not pb 同样能正确返回。题目保证一定相交，因此未对 pa 为 None 做额外处理。',
    174: '一次通过。先在本地用样例得到 9；再构造 n=1、两链表长度相同、交点位于末尾等数据，与直接查找 b 的第 n 个结点的结果比对，均一致后提交，平台评测结果为 Accepted。',
}

for idx, text in par_fills.items():
    set_para_text(doc.paragraphs[idx], text)

# submission info (table_idx, time, runtime, memory)
subs = {
    6: ('5 天前', '132 ms', '44096 kB'),
    10: ('5 天前', '24 ms', '3508 kB'),
    14: ('刚刚', '366 ms', '45336 kB'),
}
for ti, (tm, rt, mem) in subs.items():
    t = doc.tables[ti]
    set_cell_text(t.cell(0, 1), tm)
    set_cell_text(t.cell(1, 1), rt)
    set_cell_text(t.cell(1, 3), mem)

# screenshots
img_dir = r'D:\桌面\ai\imgs'
images = {
    5: img_dir + r'\codex-clipboard-9cacb0a9-82dc-4472-9056-4bf70b4f5817.jpg',
    9: img_dir + r'\codex-clipboard-d3749a31-1a17-45fd-b50b-a0b210714403.jpg',
    13: img_dir + r'\codex-clipboard-cccbdf21-fb4a-4540-b896-2496a2889e81.jpg',
}
for ti, path in images.items():
    insert_image(doc.tables[ti].cell(0, 0), path, 6.0)

# source code tables
code_tables = {7: P0010_CODE, 11: P0030_CODE, 15: P0050_CODE}
for ti, code in code_tables.items():
    set_cell_text(doc.tables[ti].cell(0, 0), numbered(code), font_name='Consolas', size=9)

# explanation tables
exp = {
    8: [
        ('1', 'import sys', '导入 sys 模块，以便用 sys.stdin.read 一次性读入全部数据，避免逐行或逐个输入导致超时。'),
        ('4', 'def main():', '定义主函数，把读入、合并、输出逻辑封装在一起，便于组织代码。'),
        ('5—7', 'data = sys.stdin.read().split()\nif not data: return', '一次性读入整个标准输入并按空白切分；读不到数据时直接返回，避免空输入访问 data[0] 越界。'),
        ('8—10', 'n = int(data[0])\na = list(map(int, data[1:1+n]))\nb = list(map(int, data[1+n:1+2*n]))', '第 0 个数字是 n；用切片分别取出两段长度为 n 的序列并转换为 int 列表。'),
        ('11—13', 'i = 0\nj = 0\nresult = []', '初始化双指针 i、j 分别指向 a、b 的当前元素，result 保存合并结果。'),
        ('14—20', 'while i < n and j < n:\n    if a[i] <= b[j]: ...\n    else: ...', '二路归并核心：两段都没取完时比较 a[i] 与 b[j]，把较小者放入 result，并后移对应指针。'),
        ('21—24', 'if i < n: result.extend(a[i:])\nif j < n: result.extend(b[j:])', '某一段先取完后，把另一段剩余的有序部分一次性追加，保证不遗漏且仍保持有序。'),
        ('25—29', 'print(" ".join(map(str, result)))\nif __name__ == "__main__": main()', '把整数结果用空格连接后一次输出，避免行尾多余空格；最后调用 main 执行。'),
    ],
    12: [
        ('1—2', 'def josephus(n, m):\n# 返回 n 只猴子...', '定义约瑟夫递推函数，注释说明该函数返回最后剩下的猴子编号。'),
        ('3', 's = 0', '设置递推基准：只有 1 个人时，最后剩下的就是 0 号（从 0 开始编号）。'),
        ('4—5', 'for i in range(2, n + 1):\n    s = (s + m) % i', '从 2 人递推到 n 人；每增加一人，把上一轮幸存位置加 m 后对当前人数取模。'),
        ('6', 'return s + 1', '把从 0 开始的编号加 1，还原成题目要求的从 1 开始的猴子编号。'),
        ('9—14', 'def main():\nresults = []\nwhile True:\n    line = input()\n    if not line.strip(): continue', '主函数循环读取多组数据，results 暂存输出；遇到空行直接跳过，避免解析错误。'),
        ('15—18', 'n, m = map(int, line.split())\nif n == 0 and m == 0: break\nresults.append(str(josephus(n, m)))', '解析每行 n、m；遇到 0 0 时结束；否则计算该组猴王编号并保存为字符串。'),
        ('19', 'print("\\n".join(results))', '所有结果用换行连接后一次输出，保证每组答案占一行。'),
        ('22—23', 'if __name__ == "__main__":\n    main()', '程序入口，调用 main 执行多组测试。'),
    ],
    16: [
        ('1—3', 'class Node:\n    def __init__(...):\n        self.data, self.next = data, next', '定义链表结点，包含数据域 data 和指针域 next，用于保存数据和指向下一个结点。'),
        ('6—13', 'class LinkList:\n    def __init__(self, lst): ...', '根据输入列表建立单链表：先创建头结点，再逐个创建新结点并串接到表尾。'),
        ('16—31', 'def genLinkList():\n    ...\n    return la, lb', '读入两个序列和 n，构造链表 la、lb；找到 a 的尾结点并把它的 next 指向 b 的第 n 个结点，形成 Y 形。'),
        ('34', 'a, b = genLinkList()', '调用 genLinkList 得到两个已交叉成 Y 字形的链表，并赋给 a、b。'),
        ('37—47', 'len_a = 0\np = a.head\n...\nlen_b = 0\np = b.head\n...', '分别遍历 a、b 求链表长度；由于 a 包含共享后缀，通常 len_a 更大。'),
        ('48—54', 'pa, pb = a.head, b.head\nif len_a > len_b: ...\nelse: ...', '让较长链表的指针先走长度差，使 pa、pb 到分叉点的剩余距离相同。'),
        ('55—58', 'while pa is not pb:\n    pa = pa.next\n    pb = pb.next\nreturn pa.data', '同步后移，直到两个指针指向同一个结点（用 is 比较对象身份），该结点就是分叉点。'),
        ('61', 'print(findJointPoint(a, b))', '题目框架调用 findJointPoint 并打印交点数据，即 b 的第 n 个结点的值。'),
    ],
}
for ti, rows in exp.items():
    t = doc.tables[ti]
    for ri, (ln, code, why) in enumerate(rows, start=1):
        set_cell_text(t.cell(ri, 0), ln, size=9)
        set_cell_text(t.cell(ri, 1), code, font_name='Consolas', size=8)
        set_cell_text(t.cell(ri, 2), why, size=9)

doc.save(DST)
print('saved', DST)
