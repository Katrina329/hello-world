# -*- coding: utf-8 -*-
import sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# (style, text) where style in: h1, h2, h3, p, b (bullet)
C = []
def h1(t): C.append(("h1", t))
def h2(t): C.append(("h2", t))
def h3(t): C.append(("h3", t))
def p(t): C.append(("p", t))
def b(t): C.append(("b", t))

h1("International Communication \u2014 Review Assignment")
p("Clermont School of Business")

h1("Exercise 1 \u2014 GreenLeaf Organics: Entering the Chinese Market")

h2("Question 1 \u2014 Cultural mismatches using Hofstede & Hall's frameworks")
p("GreenLeaf's strategy was essentially pure \u201cstandardisation\u201d: the brand transplanted its Australian marketing mix into China with only a language translation. This ignored the deep cultural distance between Australia and China. At least three specific mismatches follow.")
b("Mismatch 1 \u2014 Low-context vs. high-context communication (Hall). GreenLeaf relies on explicit, low-context, evidence-led claims (\u201cclinically proven,\u201d \u201cdermatologist tested\u201d) and assumes the message speaks for itself. China is a high-context culture in which meaning lives in relationships, implicit cues, and who delivers the message rather than in the literal words. An unsupported scientific assertion therefore feels hollow; Chinese consumers expect trust to be constructed through social proof (key opinion leaders, friends, community endorsement). This is precisely why customer feedback reported that \u201cscientific claims are met with skepticism without social proof.\u201d")
b("Mismatch 2 \u2014 Individualism vs. collectivism (Hofstede) and colour symbolism. Australia scores around 90 on individualism, whereas China scores around 20 (strongly collectivist). GreenLeaf's copy sells individual self-improvement (\u201cAnti-Aging Serum\u201d for the individual). Chinese skincare purchasing is highly social: it involves face (\u9762\u5b50), gifting, belonging and relationship, not just personal benefit. The packaging being described as \u201ccold\u201d and \u201cfuneral-like\u201d is a direct cultural mismatch: in China white is the colour of mourning, and an all-white minimalist \u201csterile\u201d aesthetic strips away the warmth and auspicious, high-status cues that Chinese consumers expect from premium skincare.")
b("Mismatch 3 \u2014 Power distance and the wrong channels (Hofstede). China scores around 80 on power distance, versus roughly 36 for Australia. High-power-distance consumers look for authority, status and official legitimacy \u2014 the \u201cofficial flagship store\u201d on Tmall or JD, endorsement by trusted experts, and recognisable brand prestige. By keeping only its own website and Instagram/Facebook (which are inaccessible or irrelevant in mainland China), GreenLeaf removed itself from the platforms where Chinese consumers look for authoritative, low-risk sellers. Consumers \u201cexpect to shop on Tmall or JD.com\u201d precisely because those platforms signal official, high-status, accountable sellers.")
p("Supporting dimensions: China's very high long-term orientation (around 87 vs. Australia's 21) means consumers expect brands to invest for the long term and build relationships; GreenLeaf's quick, transactional, standardised entry signalled the opposite.")

h2("Question 2 \u2014 Digital and payment strategy failures, and the right choices for China")
p("Failures:")
b("Platforms/channels: GreenLeaf relied on its own direct-to-consumer site and Instagram/Facebook ads \u2014 channels that are effectively inaccessible (Instagram and Facebook are blocked) or irrelevant in mainland China. It had no presence on WeChat, Xiaohongshu (Little Red Book), Tmall, JD.com or Douyin.")
b("Payment: it offered only credit card and PayPal, whereas most Chinese consumers do not hold international credit cards and pay through Alipay or WeChat Pay mobile QR payments.")
b("Logistics: there was no local fulfilment; cross-border shipping created delivery delays, customs friction and weak after-sales.")
b("Marketing/trust: it reused scientific, evidence-based claims with no local social proof (no KOL/KOC content, reviews, or livestream commerce).")
p("What GreenLeaf should adopt, and why these choices are \u201ccommunication\u201d that builds trust:")
b("Tmall (Tmall Global / flagship store) and JD.com \u2014 an official flagship store is not merely a sales channel; in a high-context, high-power-distance market it communicates legitimacy, official status and accountability (returns and after-sales). It answers \u201ccan I trust this brand?\u201d before a single product claim is read.")
b("WeChat (Official Account + Mini Program) and Xiaohongshu \u2014 WeChat is the daily relationship and CRM platform; Xiaohongshu is where skincare decisions are actually researched and where KOL/KOC content creates social proof. Being present here communicates that the brand speaks the local language and respects local social norms.")
b("Alipay and WeChat Pay \u2014 offering the payment methods people already trust communicates understanding of the local consumer and removes a decisive barrier; a familiar, seamless payment is a trust signal, not just a transaction step.")
b("Local logistics (bonded warehouse + SF Express / Cainiao, easy returns) \u2014 fast, trackable, reliable delivery communicates service quality and lowers perceived risk in a market with high sensitivity to online-shopping uncertainty.")
b("Douyin/livestream commerce and KOL/KOC seeding \u2014 livestreaming and influencer endorsement translate low-context \u201cscientific\u201d evidence into witnessed, social proof \u2014 converting claims into high-context trust.")

h2("Question 3 \u2014 Three-phase adaptation plan (Months 1\u20132, 3\u20134, 5\u20136)")
b("Phase 1 \u2014 Localise the message and establish legitimacy (Months 1\u20132): redesign packaging and naming for China \u2014 replace funeral-like white with warm, auspicious, premium cues and give products elegant, evocative Chinese names instead of blunt literal translations; open official WeChat (Official Account + Mini Program), a Tmall/JD storefront and a Xiaohongshu account; integrate Alipay and WeChat Pay and arrange bonded-warehouse cross-border logistics. Framework: Hall's high-context communication (embed the brand in local platforms and relationships) plus Hofstede's collectivism and power distance (official flagship = legitimacy; auspicious colours and elegant names = social meaning).")
b("Phase 2 \u2014 Build social proof and relationship trust (Months 3\u20134): launch KOL/KOC seeding on Xiaohongshu and Weibo, run Douyin livestream commerce with trusted hosts, and generate authentic reviews and before/after social proof; re-frame messaging from \u201cclinically proven\u201d to \u201ctrusted, recommended and loved\u201d, leading with endorsement and using science as supporting evidence; pilot a WeChat loyalty/community loop with gifting occasions. Framework: Hall's high-context communication (social proof = contextual trust) and Hofstede's collectivism + long-term orientation (relationship and face-building).")
b("Phase 3 \u2014 Scale and deepen long-term commitment (Months 5\u20136): scale paid performance marketing on the platforms proven in Phase 2 and optimise for key shopping festivals (618, Double 11); move to domestic warehousing and strengthen after-sales/returns; consider offline pop-ups or retail to reinforce premium status; introduce localised product lines and a Chinese brand story that signals long-term investment. Framework: Hofstede's long-term orientation (sustained investment signals commitment) and high-context relationship continuity (guanxi maintained over time), reinforced by power distance through high-status partnerships.")

h1("Exercise 2 \u2014 NordPay: Market Entry Communication Strategy")
p("Selected market: South Korea")

h2("Question 1 \u2014 Why South Korea is viable, the cultural risks, and a two-dimension comparison")
p("Why South Korea is viable: it has one of the world's highest smartphone-penetration rates and a digitally native population; mobile payment and fintech are already normalised (KakaoPay, Naver Pay, Toss), so NordPay does not need to educate the market on the category \u2014 only to differentiate. Strong data/privacy awareness partially aligns with NordPay's privacy-led value proposition, and a flat, transparent subscription can appeal to consumers wary of hidden fees. Sophisticated, high-income, tech-literate urban segments also match NordPay's premium, minimalist positioning.")
p("Cultural risks and a two-dimension comparison:")
b("Individualism vs. collectivism: Scandinavia is highly individualist (Sweden ~71, Norway ~69, Denmark ~74) while South Korea is strongly collectivist (~18). NordPay's core message of \u201cpersonal financial independence\u201d and individual self-reliance can feel self-centred or isolating in Korea, where financial decisions are often family- and group-oriented and where saving face and social comparison matter.")
b("Uncertainty avoidance: Scandinavia is low on uncertainty avoidance (Sweden ~29, Denmark ~23, Norway ~50) while South Korea is very high (~85). Koreans demand clear guarantees, visible and fast support, strong security and proven reliability; NordPay's email-only 24-hour support and minimalist \u201ctrust us\u201d stance are insufficient for a high-UAI market that needs immediate reassurance.")
p("Supporting dimensions: power distance (Korea ~60 vs. Scandinavia ~18\u201331) means consumers expect official recognition and clear authority cues; long-term orientation (Korea ~100) implies long-term commitment and loyalty programmes matter; and Korea is more high-context than Scandinavia's low-context directness, so communication must lean on implicit, relational and group-consensus cues.")

h2("Question 2 \u2014 Redesigning NordPay's communication approach (each element justified)")
b("Colour/symbolism: retain the trust-evoking blue base but soften the \u201ccool\u201d blue/grey palette with warmth and clearer visual hierarchy. In a high-context, high-uncertainty-avoidance market, visual cues carry meaning; an austere, overly minimal interface can read as cold and untrustworthy rather than premium. Concept: Hall's high-context communication and Hofstede's uncertainty avoidance.")
b("Messaging tone: shift from \u201cpersonal independence and data privacy\u201d (individualist framing) to collective security and reliability \u2014 \u201csafe for you and your family\u201d, \u201ctrusted by your community\u201d, \u201cclear and dependable\u201d \u2014 emphasising guarantees, security certifications and fee transparency in a reassuring way. Concept: Hofstede's individualism\u2194collectivism and uncertainty avoidance.")
b("Platform selection: replace LinkedIn and organic Twitter/X with KakaoTalk (dominant messaging, CRM and support channel), Naver (search, blog and Caf\u00e9 communities), Instagram and YouTube for mass reach, and local fintech/tech communities. Concept: high-context local ecosystems \u2014 a brand must be present where trust is socially constructed.")
b("Payment/logistics integration: integrate KakaoPay, Naver Pay, Toss and local bank transfer, and partner with at least one established local bank or fintech; distribute through Korean app stores and local settlement rails. Local integration removes friction and signals legitimacy and long-term commitment. Concept: trust as communication \u2014 local payment integration reduces uncertainty.")
b("Customer support model: move from email-only 24-hour response to KakaoTalk live chat and Korean-language support with same-day response, plus in-app FAQs and, for high-value users, human phone/video support. Concept: Hofstede's uncertainty avoidance (immediate, human reassurance) and high-context relationship orientation (personal, relational service builds trust).")

h2("Question 3 \u2014 Top three risks, mitigations, and three KPIs")
b("Risk 1 \u2014 Dominant local competition (KakaoPay, Naver Pay, Toss). Mitigation: differentiate on NordPay's unique budgeting/privacy value rather than competing head-on, and seek a strategic partnership or integration with an incumbent to gain access instead of fighting.")
b("Risk 2 \u2014 Regulatory/compliance and data-localisation requirements. Mitigation: engage local legal and compliance counsel early, comply with Korean data-protection law (PIPA) and financial regulation, commit to onshore data storage, and obtain approvals before launch.")
b("Risk 3 \u2014 Cultural/trust misalignment (high uncertainty avoidance + collectivist expectations). Mitigation: localise messaging, support and design; build social proof through Korean influencers and community endorsements; pilot with early users and iterate; signal long-term commitment with a local team and partnerships.")
p("Three KPIs for the first six months:")
b("1) Adoption/awareness \u2014 app installs and monthly active users (MAU), plus aided brand awareness or Naver search volume: measures whether communication is reaching and attracting the market.")
b("2) Activation/trust \u2014 onboarding completion rate and the share of users who link a local payment method (KakaoPay/Naver Pay): measures whether trust-building communication converts interest into usage.")
b("3) Engagement/retention \u2014 30-day retention and NPS (or CSAT): measures whether the communication and product promise translate into sustained trust and satisfaction.")

# ---------------- emit markdown ----------------
md_lines = []
for kind, t in C:
    if kind == "h1":
        md_lines.append("# " + t + "\n")
    elif kind == "h2":
        md_lines.append("## " + t + "\n")
    elif kind == "h3":
        md_lines.append("### " + t + "\n")
    elif kind == "p":
        md_lines.append(t + "\n")
    elif kind == "b":
        md_lines.append("- " + t + "\n")
md_path = r"D:\桌面\ai\homework_review_answer.md"
with open(md_path, "w", encoding="utf-8") as f:
    f.write("\n".join(md_lines))
print("MD written:", md_path)

# ---------------- emit docx ----------------
doc = Document()
# base font
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

for kind, t in C:
    if kind == "h1":
        doc.add_heading(t, level=1)
    elif kind == "h2":
        doc.add_heading(t, level=2)
    elif kind == "h3":
        doc.add_heading(t, level=3)
    elif kind == "p":
        doc.add_paragraph(t)
    elif kind == "b":
        doc.add_paragraph(t, style="List Bullet")
docx_path = r"D:\桌面\ai\homework_review_answer.docx"
doc.save(docx_path)
print("DOCX written:", docx_path)